"""Safely restore MEV02 LC identifiers from the corrected source workbook."""

from __future__ import annotations

import argparse
import re
import sqlite3
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


LC_PATTERN = re.compile(r"LC\d{2}-\d{2}")


def source_lc_rows(source: Path):
    workbook = load_workbook(source, data_only=False, read_only=False)
    sheet = max(workbook.worksheets, key=lambda item: item.max_row)
    by_name = defaultdict(list)
    by_code = defaultdict(list)
    for row_number in range(1, sheet.max_row + 1):
        code = sheet.cell(row_number, 3).value
        name = sheet.cell(row_number, 4).value
        if not isinstance(code, str) or not isinstance(name, str):
            continue
        code = code.strip()
        name = name.strip()
        if LC_PATTERN.fullmatch(code) and name:
            by_name[name].append((code, row_number))
            by_code[code].append((name, row_number))
    return sheet.title, by_name, by_code


def build_plan(connection, source: Path):
    sheet_name, by_name, by_code = source_lc_rows(source)
    project = connection.execute("SELECT id FROM projects WHERE code = 'MEV02'").fetchone()
    if not project:
        raise RuntimeError("MEV02 project was not found")
    project_id = project[0]
    auto_rows = connection.execute(
        """
        SELECT id, func_id, name
        FROM functions
        WHERE project_id = ? AND func_id LIKE 'LC-AUTO-%'
        ORDER BY func_id
        """,
        (project_id,),
    ).fetchall()
    safe = []
    conflicts = []
    for function_id, old_function_id, name in auto_rows:
        matches = by_name.get(name, [])
        if len(matches) != 1:
            conflicts.append((old_function_id, name, "source title is not unique"))
            continue
        new_function_id, source_row = matches[0]
        duplicate_rows = by_code.get(new_function_id, [])
        if len(duplicate_rows) != 1:
            detail = "; ".join(f"{item_name}@{row}" for item_name, row in duplicate_rows)
            conflicts.append((old_function_id, name, f"duplicate {new_function_id}: {detail}"))
            continue
        occupied = connection.execute(
            "SELECT id, name FROM functions WHERE project_id = ? AND func_id = ?",
            (project_id, new_function_id),
        ).fetchone()
        if occupied and occupied[0] != function_id:
            conflicts.append((old_function_id, name, f"{new_function_id} already belongs to {occupied[1]}"))
            continue
        old_requirement_id = f"SYS-{old_function_id}"
        new_requirement_id = f"SYS-{new_function_id}"
        requirement = connection.execute(
            "SELECT id, source, version FROM requirements WHERE project_id = ? AND req_id = ?",
            (project_id, old_requirement_id),
        ).fetchone()
        if not requirement:
            conflicts.append((old_function_id, name, f"missing requirement {old_requirement_id}"))
            continue
        req_occupied = connection.execute(
            "SELECT id, title FROM requirements WHERE project_id = ? AND req_id = ?",
            (project_id, new_requirement_id),
        ).fetchone()
        if req_occupied and req_occupied[0] != requirement[0]:
            conflicts.append((old_function_id, name, f"{new_requirement_id} already belongs to {req_occupied[1]}"))
            continue
        safe.append({
            "project_id": project_id,
            "function_pk": function_id,
            "old_function_id": old_function_id,
            "new_function_id": new_function_id,
            "requirement_pk": requirement[0],
            "old_requirement_id": old_requirement_id,
            "new_requirement_id": new_requirement_id,
            "old_source": requirement[1] or "",
            "old_version": requirement[2] or 1,
            "name": name,
            "source_row": source_row,
            "sheet_name": sheet_name,
        })
    return safe, conflicts


def apply_plan(connection, source: Path, safe):
    now = datetime.now().isoformat(timespec="seconds")
    source_name = source.name
    reason = f"Corrected from {source_name}; unique title and LC identifier match."
    with connection:
        for item in safe:
            provenance = f"{item['old_source']}；编号校正依据：{source_name} / {item['sheet_name']}第{item['source_row']}行"
            connection.execute(
                "UPDATE functions SET func_id = ?, updated_at = ? WHERE id = ? AND project_id = ?",
                (item["new_function_id"], now, item["function_pk"], item["project_id"]),
            )
            connection.execute(
                """
                UPDATE requirements
                SET req_id = ?, source = ?, version = ?, updated_at = ?
                WHERE id = ? AND project_id = ?
                """,
                (
                    item["new_requirement_id"], provenance, item["old_version"] + 1, now,
                    item["requirement_pk"], item["project_id"],
                ),
            )
            connection.execute(
                """
                INSERT INTO change_records
                (project_id, requirement_id, change_type, field_name, old_value, new_value,
                 reason, changed_by, changed_at)
                VALUES (?, ?, 'identifier_correction', 'req_id', ?, ?, ?, 'Codex', ?)
                """,
                (
                    item["project_id"], item["requirement_pk"], item["old_requirement_id"],
                    item["new_requirement_id"], reason, now,
                ),
            )
            connection.execute(
                """
                INSERT INTO change_records
                (project_id, requirement_id, change_type, field_name, old_value, new_value,
                 reason, changed_by, changed_at)
                VALUES (?, ?, 'source_correction', 'source', ?, ?, ?, 'Codex', ?)
                """,
                (
                    item["project_id"], item["requirement_pk"], item["old_source"],
                    provenance, reason, now,
                ),
            )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--db", type=Path, required=True)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--expect-safe", type=int)
    args = parser.parse_args()
    with sqlite3.connect(args.db) as connection:
        connection.execute("PRAGMA foreign_keys = ON")
        safe, conflicts = build_plan(connection, args.source)
        print(f"SAFE={len(safe)} CONFLICTS={len(conflicts)}")
        for item in safe:
            print(f"SAFE {item['old_function_id']} -> {item['new_function_id']} | {item['name']}")
        for old_id, name, detail in conflicts:
            print(f"HOLD {old_id} | {name} | {detail}")
        if args.expect_safe is not None and len(safe) != args.expect_safe:
            raise RuntimeError(f"Expected {args.expect_safe} safe changes, found {len(safe)}")
        if args.apply:
            apply_plan(connection, args.source, safe)
            print(f"APPLIED={len(safe)}")
        integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
        foreign_keys = connection.execute("PRAGMA foreign_key_check").fetchall()
        print(f"INTEGRITY={integrity} FOREIGN_KEYS={len(foreign_keys)}")


if __name__ == "__main__":
    main()
