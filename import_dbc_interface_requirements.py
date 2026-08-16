"""Import reviewed DBC metadata as MEV02 signals and interface requirements.

The import is idempotent. It keeps one platform Signal and one Interface
requirement per unique DBC signal name. Routed copies on other CAN buses are
preserved as rationale/evidence on the interface requirement.
"""

from __future__ import annotations

import argparse
import json
import re
import sqlite3
import zlib
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


PROJECT_CODE = "MEV02"
MODULE_CODE = "MEV02-DBC-IF"
ROOT_REQ_ID = "VEH-DBC-ROOT"
CHANGED_BY = "Codex DBC importer"
CHANGE_REASON = "依据MEV02 PTCAN/CDCAN/DCAN原始DBC创建或更新接口需求"


def norm(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def fmt(value: object) -> str:
    if value is None or value == "":
        return "—"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def now() -> str:
    return datetime.now(timezone.utc).replace(tzinfo=None).isoformat(sep=" ", timespec="seconds")


def load_signals(workbook: Path) -> dict[str, list[dict]]:
    wb = load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["DBC信号清单"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    grouped: dict[str, list[dict]] = defaultdict(list)
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        grouped[str(row["信号"])].append(row)
    return dict(grouped)


def primary(rows: list[dict]) -> dict:
    bus_order = {"PTCAN": 0, "CDCAN": 1, "DCAN": 2}
    return sorted(
        rows,
        key=lambda row: (
            str(row.get("报文") or "").startswith("Routing_"),
            bus_order.get(str(row.get("总线")), 9),
            str(row.get("报文")),
        ),
    )[0]


def enum_text(row: dict) -> str:
    return str(row.get("枚举") or "").strip()


def source_text(row: dict) -> str:
    return f"MEV02_Vehicle_CAN_Matrix_v1.8.2_0803/{Path(str(row['源文件'])).name}"


def requirement_text(row: dict) -> str:
    unit = str(row.get("单位") or "无单位")
    receivers = str(row.get("接收节点") or "DBC未显式定义")
    enum_clause = f" 枚举定义为：{enum_text(row)}。" if enum_text(row) else ""
    return (
        f"{row.get('发送节点') or '发送节点待确认'}应在{row['总线']}网络上，通过报文"
        f"{row['报文']}（ID {row['报文ID']}，DLC {fmt(row['DLC'])} Byte）以"
        f"{fmt(row['周期ms'])} ms周期向{receivers}发送信号{row['信号']}。"
        f"该信号起始位为{fmt(row['起始位'])}，长度为{fmt(row['长度'])} bit，"
        f"采用{row['字节序']}字节序和{row['数值类型']}数值类型；物理值按"
        f"Raw×{fmt(row['Factor'])}+{fmt(row['Offset'])}换算，范围为"
        f"[{fmt(row['最小值'])}, {fmt(row['最大值'])}] {unit}。{enum_clause}"
    )


def acceptance_text(row: dict) -> str:
    enum_clause = f"，并逐项验证枚举“{enum_text(row)}”" if enum_text(row) else ""
    return (
        f"使用CAN总线分析工具在{row['总线']}上采集报文{row['报文']}（{row['报文ID']}）；"
        f"验证发送周期为{fmt(row['周期ms'])} ms、DLC为{fmt(row['DLC'])} Byte、"
        f"起始位为{fmt(row['起始位'])}、长度为{fmt(row['长度'])} bit、字节序为{row['字节序']}；"
        f"对原始值进行Raw×{fmt(row['Factor'])}+{fmt(row['Offset'])}换算后，"
        f"结果应处于[{fmt(row['最小值'])}, {fmt(row['最大值'])}] {row.get('单位') or '无单位'}"
        f"{enum_clause}。所有检查项满足时判定通过。"
    )


def rationale_text(rows: list[dict], main: dict) -> str:
    routes = []
    for row in rows:
        if row is main:
            continue
        routes.append(
            f"{row['总线']}:{row['报文']}({row['报文ID']})，"
            f"{row.get('发送节点') or '发送节点未定义'}→{row.get('接收节点') or '接收节点未定义'}"
        )
    base = "由原始DBC自动生成的Draft接口需求；DBC是接口参数证据，不证明内部控制策略或安全完整性。"
    return base + (" 同名路由副本：" + "；".join(routes) if routes else "")


def req_key(signal_name: str, used: dict[str, str]) -> str:
    key = f"IF-DBC-{zlib.crc32(signal_name.encode('utf-8')) & 0xFFFFFFFF:08X}"
    if key in used and used[key] != signal_name:
        raise RuntimeError(f"接口需求编号哈希冲突：{used[key]} / {signal_name}")
    used[key] = signal_name
    return key


def upsert_requirement(conn, project_id: int, values: dict, stats: dict) -> int:
    row = conn.execute(
        "SELECT * FROM requirements WHERE project_id=? AND req_id=?",
        (project_id, values["req_id"]),
    ).fetchone()
    timestamp = now()
    fields = [
        "module_id", "parent_id", "title", "text", "req_type", "req_level", "req_nature", "status", "priority",
        "asil_level", "owner", "source", "rationale", "verification_method",
        "acceptance_criteria", "sort_order",
    ]
    if row is None:
        columns = ["project_id", "req_id", *fields, "version", "created_at", "updated_at"]
        data = [project_id, values["req_id"], *(values.get(k) for k in fields), 1, timestamp, timestamp]
        placeholders = ",".join("?" for _ in columns)
        conn.execute(
            f"INSERT INTO requirements ({','.join(columns)}) VALUES ({placeholders})",
            data,
        )
        requirement_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        conn.execute(
            "INSERT INTO change_records (project_id,requirement_id,change_type,field_name,old_value,new_value,reason,changed_by,changed_at) "
            "VALUES (?,?,?,?,?,?,?,?,?)",
            (project_id, requirement_id, "Create", "requirement", "", values["text"], CHANGE_REASON, CHANGED_BY, timestamp),
        )
        stats["requirements_created"] += 1
        return requirement_id

    requirement_id = row["id"]
    changes = [(field, row[field], values.get(field)) for field in fields if row[field] != values.get(field)]
    if changes:
        for field, old, new in changes:
            conn.execute(
                "INSERT INTO change_records (project_id,requirement_id,change_type,field_name,old_value,new_value,reason,changed_by,changed_at) "
                "VALUES (?,?,?,?,?,?,?,?,?)",
                (project_id, requirement_id, "Update", field, "" if old is None else str(old), "" if new is None else str(new), CHANGE_REASON, CHANGED_BY, timestamp),
            )
            conn.execute(f"UPDATE requirements SET {field}=? WHERE id=?", (new, requirement_id))
        conn.execute("UPDATE requirements SET version=version+1, updated_at=? WHERE id=?", (timestamp, requirement_id))
        stats["requirements_updated"] += 1
    else:
        stats["requirements_unchanged"] += 1
    return requirement_id


def run(database: Path, workbook: Path) -> dict:
    grouped = load_signals(workbook)
    stats = defaultdict(int)
    conn = sqlite3.connect(database)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    try:
        with conn:
            project = conn.execute("SELECT id FROM projects WHERE code=?", (PROJECT_CODE,)).fetchone()
            if project is None:
                raise RuntimeError("未找到MEV02项目")
            project_id = project["id"]
            timestamp = now()
            module = conn.execute(
                "SELECT id FROM requirement_modules WHERE project_id=? AND code=?",
                (project_id, MODULE_CODE),
            ).fetchone()
            if module is None:
                conn.execute(
                    "INSERT INTO requirement_modules (project_id,code,title,description,status,owner,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?)",
                    (project_id, MODULE_CODE, "MEV02 DBC接口需求规范", "由PTCAN、CDCAN、DCAN原始DBC生成的结构化接口需求", "Draft", "整车网络/系统架构", timestamp, timestamp),
                )
                module_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
                stats["modules_created"] += 1
            else:
                module_id = module["id"]

            root_values = {
                "req_id": ROOT_REQ_ID, "module_id": module_id, "parent_id": None,
                "title": "MEV02整车CAN通信接口",
                "text": "MEV02车辆各ECU应按照受控的PTCAN、CDCAN和DCAN原始DBC定义交换CAN网络数据；具体报文及信号约束由本规范下级接口需求规定。",
                "req_type": "Vehicle", "req_level": "Vehicle", "req_nature": "Interface",
                "status": "Draft", "priority": "Medium", "asil_level": None,
                "owner": "整车网络/系统架构", "source": "PTCAN.dbc；CDCAN.dbc；DCAN.dbc",
                "rationale": "作为原始DBC接口需求的规范根节点，不用于推断ASIL或内部控制逻辑。",
                "verification_method": "Analysis",
                "acceptance_criteria": "本规范中每条接口需求均应分配到唯一平台信号，并可追溯到原始DBC文件。",
                "sort_order": 0,
            }
            root_id = upsert_requirement(conn, project_id, root_values, stats)

            db_signals = conn.execute("SELECT * FROM signals WHERE project_id=?", (project_id,)).fetchall()
            exact = {row["name"]: row for row in db_signals}
            normalized = defaultdict(list)
            for row in db_signals:
                normalized[norm(row["name"])].append(row)

            used_keys = {}
            for order, signal_name in enumerate(sorted(grouped), 1):
                rows = grouped[signal_name]
                main = primary(rows)
                signal_row = exact.get(signal_name)
                if signal_row is None:
                    matches = normalized.get(norm(signal_name), [])
                    signal_row = matches[0] if len(matches) == 1 else None
                description = (
                    f"原始DBC接口数据\n总线: {main['总线']}\n发送节点: {main.get('发送节点') or '未定义'}\n"
                    f"接收节点: {main.get('接收节点') or '未定义'}\n枚举: {enum_text(main) or '无'}\n"
                    f"来源: {source_text(main)}\n{rationale_text(rows, main)}"
                )
                signal_values = (
                    signal_name, signal_name, main["报文"], main["报文ID"], main["DLC"], main["起始位"],
                    main["长度"], main["Factor"], main["Offset"], main["最小值"], main["最大值"],
                    main.get("单位") or None, main["字节序"], main["数值类型"], int(main["周期ms"]),
                    main.get("发送节点") or None, description,
                )
                if signal_row is None:
                    conn.execute(
                        "INSERT INTO signals (project_id,name,signal_id,message_name,message_id,dlc,start_bit,bit_length,factor,offset,min_value,max_value,unit,byte_order,value_type,cycle_time,module,description,created_at) "
                        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                        (project_id, *signal_values, timestamp),
                    )
                    signal_db_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
                    stats["signals_created"] += 1
                else:
                    signal_db_id = signal_row["id"]
                    conn.execute(
                        "UPDATE signals SET name=?,signal_id=?,message_name=?,message_id=?,dlc=?,start_bit=?,bit_length=?,factor=?,offset=?,min_value=?,max_value=?,unit=?,byte_order=?,value_type=?,cycle_time=?,module=?,description=? WHERE id=? AND project_id=?",
                        (*signal_values, signal_db_id, project_id),
                    )
                    stats["signals_updated"] += 1

                req_values = {
                    "req_id": req_key(signal_name, used_keys), "module_id": module_id, "parent_id": root_id,
                    "title": f"接口信号：{signal_name}", "text": requirement_text(main),
                    "req_type": "Interface", "req_level": "Interface", "req_nature": "Interface",
                    "status": "Draft", "priority": "Medium", "asil_level": None,
                    "owner": main.get("发送节点") or "整车网络/系统架构", "source": source_text(main),
                    "rationale": rationale_text(rows, main), "verification_method": "Test",
                    "acceptance_criteria": acceptance_text(main), "sort_order": order,
                }
                interface_id = upsert_requirement(conn, project_id, req_values, stats)
                conn.execute(
                    "INSERT OR IGNORE INTO trace_links (project_id,source_type,source_id,target_type,target_id,link_type,description,created_by,created_at) VALUES (?,?,?,?,?,?,?,?,?)",
                    (project_id, "Requirement", root_id, "Requirement", interface_id, "decomposes", "整车CAN通信接口分解为DBC信号接口需求", CHANGED_BY, timestamp),
                )
                stats["root_links"] += conn.execute("SELECT changes()").fetchone()[0]
                conn.execute(
                    "INSERT OR IGNORE INTO trace_links (project_id,source_type,source_id,target_type,target_id,link_type,description,created_by,created_at) VALUES (?,?,?,?,?,?,?,?,?)",
                    (project_id, "Requirement", interface_id, "Signal", signal_db_id, "allocates", f"接口需求分配到原始DBC信号 {signal_name}", CHANGED_BY, timestamp),
                )
                stats["signal_links"] += conn.execute("SELECT changes()").fetchone()[0]

            stats["dbc_unique_signals"] = len(grouped)
            stats["interface_requirements"] = conn.execute(
                "SELECT count(*) FROM requirements WHERE project_id=? AND module_id=? AND req_type='Interface'",
                (project_id, module_id),
            ).fetchone()[0]
            stats["total_requirements"] = conn.execute(
                "SELECT count(*) FROM requirements WHERE project_id=?", (project_id,)
            ).fetchone()[0]
            stats["total_trace_links"] = conn.execute(
                "SELECT count(*) FROM trace_links WHERE project_id=?", (project_id,)
            ).fetchone()[0]
    finally:
        conn.close()
    return dict(stats)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("database", type=Path)
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    print(json.dumps(run(args.database.resolve(), args.workbook.resolve()), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
